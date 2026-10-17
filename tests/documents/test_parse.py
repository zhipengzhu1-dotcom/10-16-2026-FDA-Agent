import io
import zipfile
from itertools import pairwise
from pathlib import Path

import docx
import pymupdf
import pytest
from docx.oxml import parse_xml

from fda_copilot.cli import main
from fda_copilot.documents.parse import ParsedDocument, UnsupportedDocument, parse_document

FIXTURES = Path(__file__).parent / "fixtures"
DOCX = FIXTURES / "protocol.docx"
PDF = FIXTURES / "protocol.pdf"


def parse(path: Path) -> ParsedDocument:
    return parse_document(path.name, path.read_bytes())


W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def docx_with_body_xml(*blocks: str) -> bytes:
    document = docx.Document()
    document.add_heading("1. Consent", level=1)
    document.add_paragraph("before")
    for block in blocks:
        document.element.body.sectPr.addprevious(parse_xml(block))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pdf_with_lines(lines: list[tuple[str, float]], toc: list[list[object]] | None = None) -> bytes:
    with pymupdf.open() as document:
        page = document.new_page()
        for y, (text, size) in enumerate(lines, start=1):
            page.insert_text((72, 30 * y), text, fontname="helv", fontsize=size)
        if toc is not None:
            document.set_toc(toc)
        return document.tobytes()


def outlined_pdf() -> bytes:
    with pymupdf.open(PDF) as document:
        document.set_toc(
            [[1, "Introduction", 1], [1, "Informed Consent and Participant Rights", 2]]
        )
        return document.tobytes()


def test_docx_headings() -> None:
    assert [section.heading for section in parse(DOCX).sections] == [
        "",
        "Example Protocol EX-001",
        "1. Introduction",
        "1.1 Background",
        "2. Informed Consent",
        "3. Safety Reporting",
    ]


def test_docx_section_text_includes_table_rows() -> None:
    consent = parse(DOCX).sections[4]
    assert consent.text == (
        "2. Informed Consent\nRisk\tLevel\nHeadache\tLow\n"
        "Each participant signs the form before dosing.\n"
    )


def test_pdf_font_size_headings() -> None:
    assert [section.heading for section in parse(PDF).sections] == [
        "",
        "Example Protocol EX-001",
        "Contents",
        "1. Introduction",
        "1.1 Background",
        "2. Informed Consent and Participant Rights",
        "3. Safety Reporting",
    ]


def test_pdf_outline_headings_win_over_font_sizes() -> None:
    sections = parse_document("outlined.pdf", outlined_pdf()).sections
    assert [section.heading for section in sections] == [
        "",
        "Introduction",
        "Informed Consent and Participant Rights",
    ]
    assert sections[1].text.startswith("1. Introduction\nThis study tests drug X")


@pytest.mark.parametrize(
    ("document", "index", "span", "text"),
    [
        (parse(DOCX), 2, (48, 105), "1. Introduction\nThis study tests drug X against placebo.\n"),
        (
            parse(PDF),
            3,
            (84, 169),
            "1. Introduction\nThis study tests drug X against placebo.\n"
            "Body text for this section.\n",
        ),
        (
            parse_document("outlined.pdf", outlined_pdf()),
            0,
            (0, 84),
            "Sponsor: Example Pharma\nExample Protocol EX-001\nContents\n"
            "1. Introduction ........ 1\n",
        ),
    ],
    ids=["docx", "pdf-font", "pdf-outline"],
)
def test_spans_slice_back_and_tile_the_text(
    document: ParsedDocument, index: int, span: tuple[int, int], text: str
) -> None:
    section = document.sections[index]
    assert (section.span.start, section.span.end) == span
    assert section.text == text
    spans = [(section.span.start, section.span.end) for section in document.sections]
    assert spans[0][0] == 0
    assert spans[-1][1] == len(document.text)
    assert all(end == next_start for (_, end), (next_start, _) in pairwise(spans))
    for section in document.sections:
        assert section.text == document.text[section.span.start : section.span.end]
        assert section.text


def test_parsed_document_refuses_spans_that_leave_a_gap() -> None:
    first = parse(DOCX).sections[0]
    with pytest.raises(ValueError, match="cover the whole text"):
        ParsedDocument(text=first.text + "orphan\n", sections=(first,))


def test_section_ids_are_unique() -> None:
    sections = parse(PDF).sections
    assert len({section.id for section in sections}) == len(sections)


@pytest.mark.parametrize(
    ("filename", "data"),
    [("notes.txt", b"hello"), ("scan.pdf", b"not a pdf"), ("protocol.docx", b"not a zip")],
)
def test_rejects_other_files(filename: str, data: bytes) -> None:
    with pytest.raises(UnsupportedDocument):
        parse_document(filename, data)


def test_cli_prints_one_line_per_section(capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["parse", str(DOCX)]) == 0
    assert capsys.readouterr().out == (
        "0-24\t\n"
        "24-48\tExample Protocol EX-001\n"
        "48-105\t1. Introduction\n"
        "105-151\t1.1 Background\n"
        "151-242\t2. Informed Consent\n"
        "242-309\t3. Safety Reporting\n"
    )


def test_cli_reports_unsupported_file(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    notes = tmp_path / "notes.txt"
    notes.write_text("hello")
    assert main(["parse", str(notes)]) == 1
    assert capsys.readouterr().err == (
        "fda-copilot parse: notes.txt: only .docx and .pdf are supported\n"
    )


def test_docx_includes_content_controls_in_document_order() -> None:
    data = docx_with_body_xml(
        f"<w:sdt {W}><w:sdtPr/><w:sdtContent>"
        "<w:p><w:r><w:t>inside sdt</w:t></w:r></w:p>"
        "<w:sdt><w:sdtContent><w:tbl><w:tr><w:tc><w:p><w:r><w:t>nested</w:t></w:r></w:p></w:tc>"
        "<w:tc><w:p><w:r><w:t>cell</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:sdtContent></w:sdt>"
        "</w:sdtContent></w:sdt>",
        f"<w:p {W}><w:r><w:t>after</w:t></w:r></w:p>",
    )
    assert parse_document("sdt.docx", data).text == (
        "1. Consent\nbefore\ninside sdt\nnested\tcell\nafter\n"
    )


def test_docx_shows_tracked_changes_as_accepted() -> None:
    data = docx_with_body_xml(
        f"<w:p {W}><w:r><w:t xml:space='preserve'>keep </w:t></w:r>"
        "<w:ins w:id='1' w:author='a'><w:r><w:t>inserted</w:t></w:r></w:ins>"
        "<w:del w:id='2' w:author='a'><w:r><w:delText> deleted</w:delText></w:r></w:del></w:p>",
        f"<w:tbl {W}><w:tr><w:tc><w:p><w:ins w:id='3' w:author='a'><w:r><w:t>new cell</w:t></w:r>"
        "</w:ins></w:p></w:tc></w:tr></w:tbl>",
    )
    assert parse_document("tracked.docx", data).text == (
        "1. Consent\nbefore\nkeep inserted\nnew cell\n"
    )


def malformed_docx() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(DOCX) as source, zipfile.ZipFile(buffer, "w") as target:
        for item in source.infolist():
            data = source.read(item)
            target.writestr(
                item, b"<w:document><w:body>" if item.filename == "word/document.xml" else data
            )
    return buffer.getvalue()


def test_cli_reports_malformed_docx_xml(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    broken = tmp_path / "broken.docx"
    broken.write_bytes(malformed_docx())
    assert main(["parse", str(broken)]) == 1
    assert capsys.readouterr().err == "fda-copilot parse: broken.docx: not a readable .docx file\n"


@pytest.mark.parametrize("decoy", ["Study", "Risk"])
def test_pdf_outline_skips_a_body_line_that_only_prefixes_the_title(decoy: str) -> None:
    title = {"Study": "Study Design and Rationale", "Risk": "Risk Assessment"}[decoy]
    data = pdf_with_lines(
        [("Intro text.", 11), (decoy, 11), (f"5. {title}", 14), ("Body.", 11)],
        toc=[[1, title, 1]],
    )
    sections = parse_document("decoy.pdf", data).sections
    assert [(section.heading, section.text) for section in sections] == [
        ("", f"Intro text.\n{decoy}\n"),
        (title, f"5. {title}\nBody.\n"),
    ]


@pytest.mark.parametrize(
    ("filename", "data", "message"),
    [
        ("empty.pdf", b"%PDF-1.7\n1 0 obj garbage\n", "empty.pdf: the PDF has no pages"),
        (
            "scan.pdf",
            pdf_with_lines([]),
            "scan.pdf: the PDF has no text layer; it may be a scan that needs OCR",
        ),
    ],
    ids=["no-pages", "no-text"],
)
def test_cli_rejects_pdf_without_text(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], filename: str, data: bytes, message: str
) -> None:
    path = tmp_path / filename
    path.write_bytes(data)
    assert main(["parse", str(path)]) == 1
    assert capsys.readouterr().err == f"fda-copilot parse: {message}\n"


def test_pdf_heading_drops_a_superscript_footnote_marker() -> None:
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "APPENDIX 1. EVENTS", fontname="hebo", fontsize=14)
        width = pymupdf.get_text_length("APPENDIX 1. EVENTS", fontname="hebo", fontsize=14)
        page.insert_text((72 + width, 66), "1", fontname="hebo", fontsize=8)
        page.insert_text((72, 100), "Visits happen weekly.", fontname="helv", fontsize=11)
        data = document.tobytes()
    sections = parse_document("appendix.pdf", data).sections
    assert [(section.heading, section.text) for section in sections] == [
        ("APPENDIX 1. EVENTS", "APPENDIX 1. EVENTS1\nVisits happen weekly.\n"),
    ]


def row(*cells: str) -> str:
    return "<w:tr>" + "".join(cells) + "</w:tr>"


def cell(text: str) -> str:
    return f"<w:tc><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:tc>"


def sdt(content: str) -> str:
    return f"<w:sdt><w:sdtPr/><w:sdtContent>{content}</w:sdtContent></w:sdt>"


def test_docx_keeps_rows_and_cells_wrapped_in_content_controls() -> None:
    data = docx_with_body_xml(
        f"<w:p {W}><w:r><w:t>Rows</w:t></w:r></w:p>",
        f"<w:tbl {W}>"
        + row(cell("R1A"), cell("R1B"))
        + sdt(row(cell("ROW_SDT_A"), cell("ROW_SDT_B")))
        + row(cell("R3A"), sdt(cell("CELL_LEVEL_SDT")))
        + "<w:customXml w:element='x'>"
        + row(f"<w:customXml w:element='y'>{cell('XML_A')}</w:customXml>", cell("XML_B"))
        + "</w:customXml></w:tbl>",
    )
    assert parse_document("rows.docx", data).text == (
        "1. Consent\nbefore\nRows\nR1A\tR1B\nROW_SDT_A\tROW_SDT_B\nR3A\tCELL_LEVEL_SDT\n"
        "XML_A\tXML_B\n"
    )


def test_docx_renders_a_nested_table_as_cell_lines() -> None:
    nested = (
        "<w:tbl>" + row(cell("N1A"), cell("N1B")) + sdt(row(cell("N2A"), cell("N2B"))) + "</w:tbl>"
    )
    data = docx_with_body_xml(
        f"<w:tbl {W}>"
        + row(
            cell("OUTER_A"),
            "<w:tc><w:p><w:r><w:t>P1</w:t></w:r></w:p>"
            + sdt(nested)
            + "<w:p><w:r><w:t>P2</w:t></w:r></w:p></w:tc>",
        )
        + "</w:tbl>",
    )
    assert parse_document("nested.docx", data).text == (
        "1. Consent\nbefore\nOUTER_A\tP1\nN1A\tN1B\nN2A\tN2B\nP2\n"
    )


def test_docx_reads_merged_cell_text_once() -> None:
    data = docx_with_body_xml(
        f"<w:tbl {W}>"
        + row(
            cell("SPAN").replace("<w:tc>", "<w:tc><w:tcPr><w:gridSpan w:val='2'/></w:tcPr>"),
            cell("TOP").replace("<w:tc>", "<w:tc><w:tcPr><w:vMerge w:val='restart'/></w:tcPr>"),
        )
        + row(cell("A"), cell("B"), "<w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p/></w:tc>")
        + "</w:tbl>",
    )
    assert parse_document("merged.docx", data).text == ("1. Consent\nbefore\nSPAN\tTOP\nA\tB\t\n")
