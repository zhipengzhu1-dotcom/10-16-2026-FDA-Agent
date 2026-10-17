"""Rebuild protocol.docx and protocol.pdf. Run from the repo root with `uv run python`."""

from pathlib import Path

import docx
import pymupdf

HERE = Path(__file__).parent

BODY = "Body text for this section."


def make_docx(path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Sponsor: Example Pharma")
    document.add_paragraph("Example Protocol EX-001", style="Title")
    document.add_heading("1. Introduction", level=1)
    document.add_paragraph("This study tests drug X against placebo.")
    document.add_heading("1.1 Background", level=2)
    document.add_paragraph("Drug X lowers blood pressure.")
    document.add_heading("", level=2)
    document.add_heading("2. Informed Consent", level=1)
    table = document.add_table(rows=2, cols=2)
    for row, cells in zip(table.rows, [("Risk", "Level"), ("Headache", "Low")], strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    document.add_paragraph("Each participant signs the form before dosing.")
    document.add_heading("3. Safety Reporting", level=1)
    document.add_paragraph("Report serious adverse events within 24 hours.")
    document.save(str(path))


def make_pdf(path: Path) -> None:
    document = pymupdf.open()
    pages: list[list[tuple[str, str, float]]] = [
        [
            ("helv", "Sponsor: Example Pharma", 9),
            ("hebo", "Example Protocol EX-001", 18),
            ("hebo", "Contents", 14),
            ("hebo", "1. Introduction ........ 1", 11),
            ("hebo", "1. Introduction", 14),
            ("helv", "This study tests drug X against placebo.", 11),
            ("helv", BODY, 11),
            ("hebo", "1.1 Background", 11),
            ("helv", "Drug X lowers blood pressure.", 11),
            ("hebo", "7 Days", 11),
            ("helv", BODY, 11),
            ("helv", BODY, 11),
        ],
        [
            ("hebo", "2. Informed Consent and", 14),
            ("hebo", "Participant Rights", 14),
            ("helv", "Each participant signs the form before dosing.", 11),
            ("helv", BODY, 11),
            ("hebo", "3. Safety Reporting", 14),
            ("helv", "Report serious adverse events within 24 hours.", 11),
            ("helv", BODY, 11),
        ],
    ]
    for lines in pages:
        page = document.new_page()
        y = 72.0
        for font, text, size in lines:
            page.insert_text((72, y), text, fontname=font, fontsize=size)
            y += size * 1.8
    document.save(str(path), garbage=4, deflate=True)


if __name__ == "__main__":
    make_docx(HERE / "protocol.docx")
    make_pdf(HERE / "protocol.pdf")
