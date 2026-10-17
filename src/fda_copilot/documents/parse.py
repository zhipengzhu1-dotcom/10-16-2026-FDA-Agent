import hashlib
import io
import re
import zipfile
from collections import Counter
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import PurePath
from typing import Any

import docx
import pymupdf
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph

from fda_copilot.domain.models import CharSpan, DocSection, DocSectionId


class UnsupportedDocument(ValueError):
    pass


@dataclass(frozen=True)
class ParsedDocument:
    """The parse of one uploaded file.

    `text` is the canonical extracted text: every extracted line in document order, each
    followed by one newline. A PDF line is one pymupdf text line, a DOCX line is one body
    paragraph or one table row with cells joined by tabs. Every section's span indexes into
    `text`, its `text` equals that slice, and the spans tile `text` in order with no gap and
    no overlap. Text before the first heading is a section whose heading is empty.
    """

    text: str
    sections: tuple[DocSection, ...]


@dataclass(frozen=True)
class _Line:
    text: str
    heading: str | None = None


@dataclass(frozen=True)
class _PdfLine:
    text: str
    page: int
    style: tuple[str, float, bool] | None

    @property
    def size(self) -> float:
        return self.style[1] if self.style else 0.0


_SECTION_NUMBER = re.compile(r"^\s*\d+(?:\.\d+)*\.?(?:\s|$)")
_BARE_NUMBER = re.compile(r"^\s*\d+(?:\.\d+)*\.?\s*$")
_DOTTED_NUMBER = re.compile(r"^\s*\d+\.(?:\d+\.?)*(?:\s|$)")
_TOC_LEADER = re.compile(r"\.{5,}|…{2,}")
_MAX_HEADING_CHARS = 200


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".pdf":
        lines = _pdf_lines(filename, data)
    elif suffix == ".docx":
        lines = _docx_lines(filename, data)
    else:
        raise UnsupportedDocument(f"{filename}: only .docx and .pdf are supported")
    return _assemble(lines, hashlib.sha256(data).hexdigest()[:12])


def _assemble(lines: Sequence[_Line], doc_key: str) -> ParsedDocument:
    text = "".join(f"{line.text}\n" for line in lines)
    if not text:
        return ParsedDocument(text="", sections=())
    starts: list[tuple[int, str]] = []
    offset = 0
    for line in lines:
        if line.heading is not None:
            starts.append((offset, line.heading))
        offset += len(line.text) + 1
    if not starts or starts[0][0] > 0:
        starts.insert(0, (0, ""))
    ends = [start for start, _ in starts[1:]] + [len(text)]
    sections = tuple(
        DocSection(
            id=DocSectionId(f"{doc_key}-{index:04d}"),
            heading=heading,
            text=text[start:end],
            span=CharSpan(start=start, end=end),
        )
        for index, ((start, heading), end) in enumerate(zip(starts, ends, strict=True))
    )
    return ParsedDocument(text=text, sections=sections)


def _docx_lines(filename: str, data: bytes) -> list[_Line]:
    try:
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, PackageNotFoundError, KeyError, ValueError) as error:
        raise UnsupportedDocument(f"{filename}: not a readable .docx file") from error
    lines: list[_Line] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            lines.append(_Line(block.text, _docx_heading(block)))
        elif isinstance(block, Table):
            lines.extend(_Line("\t".join(cell.text for cell in row.cells)) for row in block.rows)
    return lines


def _docx_heading(paragraph: Paragraph) -> str | None:
    style = paragraph.style.name if paragraph.style is not None else None
    heading = _squash(paragraph.text)
    if style and heading and (style == "Title" or style.startswith("Heading")):
        return heading
    return None


def _pdf_lines(filename: str, data: bytes) -> list[_Line]:
    try:
        document = pymupdf.Document(stream=data, filetype="pdf")  # type: ignore[no-untyped-call]
    except pymupdf.FileDataError as error:
        raise UnsupportedDocument(f"{filename}: not a readable .pdf file") from error
    with document:
        if document.needs_pass:
            raise UnsupportedDocument(f"{filename}: the PDF is encrypted")
        lines: list[_PdfLine] = []
        for page_number, page in enumerate(document.pages(), start=1):
            page_text: dict[str, Any] = page.get_text(  # type: ignore[no-untyped-call]
                "dict", flags=pymupdf.TEXTFLAGS_TEXT
            )
            lines.extend(
                _pdf_line(line, page_number)
                for block in page_text["blocks"]
                for line in block["lines"]
            )
        toc: list[list[Any]] = document.get_toc()
    headings = _outline_headings(toc, lines) or _font_headings(lines)
    return [_Line(line.text, headings.get(index)) for index, line in enumerate(lines)]


def _pdf_line(line: dict[str, Any], page: int) -> _PdfLine:
    spans: list[dict[str, Any]] = line["spans"]
    inked = [span for span in spans if any(char.isalnum() for char in span["text"])]
    style = None
    if inked:
        bold = all(
            span["flags"] & pymupdf.TEXT_FONT_BOLD or "bold" in span["font"].lower()
            for span in inked
        )
        style = (str(inked[0]["font"]), round(max(float(s["size"]) for s in inked), 1), bold)
    return _PdfLine("".join(span["text"] for span in spans), page, style)


def _outline_headings(toc: Sequence[Sequence[Any]], lines: Sequence[_PdfLine]) -> dict[int, str]:
    headings: dict[int, str] = {}
    cursor = 0
    for _level, title, page, *_ in toc:
        wanted = _title_key(str(title))
        if not wanted:
            continue
        for index in range(cursor, len(lines)):
            line = lines[index]
            if line.page < page:
                continue
            if line.page > page + 1:
                break
            got = _title_key(line.text)
            if got == wanted or (len(got) >= 4 and wanted.startswith(got)):
                if index > cursor and _BARE_NUMBER.match(lines[index - 1].text):
                    index -= 1
                headings[index] = _squash(str(title))
                cursor = index + 1
                break
    return headings


def _title_key(text: str) -> str:
    return _squash(_SECTION_NUMBER.sub("", text)).casefold()


def _font_headings(lines: Sequence[_PdfLine]) -> dict[int, str]:
    sizes: Counter[float] = Counter()
    for line in lines:
        if line.style:
            sizes[line.size] += len(line.text.strip())
    if not sizes:
        return {}
    body_size = sizes.most_common(1)[0][0]

    def starts_heading(line: _PdfLine) -> bool:
        if line.style is None:
            return False
        bold_number = line.style[2] and _DOTTED_NUMBER.match(line.text) is not None
        return line.size >= body_size + 0.5 or bold_number

    headings: dict[int, str] = {}
    index = 0
    while index < len(lines):
        first = lines[index]
        if not starts_heading(first):
            index += 1
            continue
        end = index + 1
        while (
            end < len(lines)
            and lines[end].style == first.style
            and lines[end].page == first.page
            and not _SECTION_NUMBER.match(lines[end].text)
        ):
            end += 1
        heading = _squash(" ".join(line.text for line in lines[index:end]))
        if (
            len(heading) <= _MAX_HEADING_CHARS
            and any(char.isalpha() for char in heading)
            and not _TOC_LEADER.search(heading)
        ):
            headings[index] = heading
        index = end
    return headings


def _squash(text: str) -> str:
    return " ".join(text.split())
